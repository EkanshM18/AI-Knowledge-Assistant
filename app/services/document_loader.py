from __future__ import annotations

import csv
import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from llama_index.core import Document
from pypdf import PdfReader


class DocumentLoader:
    supported_extensions = {".pdf", ".txt", ".docx", ".md", ".csv", ".json", ".html", ".htm"}

    def load_from_bytes(
        self,
        *,
        filename: str,
        content: bytes,
        tags: list[str] | None = None,
        source_path: str = "",
    ) -> list[Document]:
        suffix = Path(filename).suffix.lower()
        if suffix not in self.supported_extensions:
            raise ValueError(f"Unsupported file type: {suffix}")

        tags = tags or []
        timestamp = datetime.now(timezone.utc).isoformat()

        if suffix == ".pdf":
            return self._load_pdf_bytes(filename=filename, content=content, tags=tags, timestamp=timestamp, source_path=source_path)

        if suffix == ".docx":
            text = self._load_docx_text(BytesIO(content))
        elif suffix in {".html", ".htm"}:
            text = self._load_html_text(content)
        elif suffix == ".csv":
            text = self._load_csv_text(content)
        elif suffix == ".json":
            text = self._load_json_text(content)
        else:
            text = content.decode("utf-8", errors="ignore")

        return [self._build_document(filename, text, tags, timestamp, source_path=source_path)]

    def load(self, file_path: Path, tags: list[str] | None = None) -> list[Document]:
        suffix = file_path.suffix.lower()
        if suffix not in self.supported_extensions:
            raise ValueError(f"Unsupported file type: {suffix}")

        tags = tags or []
        timestamp = datetime.now(timezone.utc).isoformat()
        content = file_path.read_bytes()
        return self.load_from_bytes(
            filename=file_path.name,
            content=content,
            tags=tags,
            source_path=str(file_path),
        )

    def _load_pdf_bytes(
        self,
        *,
        filename: str,
        content: bytes,
        tags: list[str],
        timestamp: str,
        source_path: str,
    ) -> list[Document]:
        documents: list[Document] = []
        reader = PdfReader(BytesIO(content))
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if not text:
                continue
            documents.append(
                self._build_document(
                    filename,
                    text,
                    tags,
                    timestamp,
                    page_number=page_number,
                    source_path=source_path,
                )
            )
        return documents

    def _load_docx_text(self, file_like: Any) -> str:
        document = DocxDocument(file_like)
        return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())

    def _load_html_text(self, content: bytes) -> str:
        html = content.decode("utf-8", errors="ignore")
        soup = BeautifulSoup(html, "lxml")
        return soup.get_text(separator="\n", strip=True)

    def _load_csv_text(self, content: bytes) -> str:
        rows: list[str] = []
        decoded = content.decode("utf-8", errors="ignore")
        reader = csv.reader(decoded.splitlines())
        for row in reader:
            rows.append(" | ".join(cell.strip() for cell in row))
        return "\n".join(rows)

    def _load_json_text(self, content: bytes) -> str:
        decoded = content.decode("utf-8", errors="ignore")
        payload = json.loads(decoded)
        return json.dumps(payload, indent=2, ensure_ascii=True)

    def _build_document(
        self,
        filename: str,
        text: str,
        tags: list[str],
        timestamp: str,
        page_number: int | None = None,
        source_path: str = "",
    ) -> Document:
        metadata = {
            "filename": filename,
            "page_number": page_number,
            "source": source_path or filename,
            "tags": tags,
            "timestamp": timestamp,
        }
        doc_id = f"{filename}:{page_number or 0}"
        return Document(text=text, metadata=metadata, doc_id=doc_id)
